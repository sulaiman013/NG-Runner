"""
Saving the IC Conversion workbook, and reading an .xlsx back with the standard library to check it.

Why openpyxl for saving: it is what Jim's toolkit writes every workbook with, and BUDGET.md settles
how ("Ship the openpyxl file", fullCalcOnLoad set, so Excel computes the totals on opening). It is
imported inside save() only, so this module loads in NG Runner's own venv, which has no openpyxl and
mustn't (pyproject.toml). Saving runs where openpyxl is: the tools' Python and Jim's container.

Why a reader of its own: so the saved file is checked in Python everywhere, with no LibreOffice (Jim's
PC has none) and no openpyxl, and so his workbooks can be read the same way to be scored. It reads
what a workbook of this kind holds: shared and inline strings, numbers, booleans, formulas (as their
text: openpyxl caches no value, and the evaluator in sheets.py computes them), bold and font colour,
and column widths. A file it can't read raises WorkbookUnreadable, in plain words.

check() and check_book() say, in plain words, every way a workbook doesn't say what its conversion
says: a sheet or a cell not as laid out, a total that doesn't add up, a formula that can't be
evaluated, a part row more or fewer than the conversion's parts, a bucket's sheets where "could not
verify" belongs, or a line that could not be verified missing from its list. An empty list is clean.
"""
from __future__ import annotations

import os
import re
import zipfile
from collections.abc import Mapping
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree

from tools.workbook.sheets import (COULD_NOT_VERIFY, PARTS, PULLED, TAKEOFF, UNPLACED_HEADER, Book, Cell,
                                   Formula, Sheet, Unevaluable, column_letter, column_number, evaluate,
                                   layout, named, unplaced_words)

_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL = re.compile(r"^(?P<col>[A-Z]{1,3})(?P<row>\d+)$")
SQFT_TOLERANCE = 0.05          # as tools/score_takeoff.py
TAPE_TOLERANCE = 0.2


class WorkbookUnreadable(Exception):
    """An .xlsx that could not be read, in plain words."""


# ----- saving (openpyxl)


def save(book: Book, path: str | os.PathLike[str]) -> None:
    """The book as an .xlsx at path, written by openpyxl with fullCalcOnLoad set."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.workbook.properties import CalcProperties

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in book.sheets:
        ws = wb.create_sheet(sheet.title)
        for (row, col), cell in sorted(sheet.cells.items()):
            if cell.value is None and not cell.bold:
                continue
            target = ws.cell(row, col, cell.value.text if isinstance(cell.value, Formula) else cell.value)
            if cell.bold or cell.colour:
                target.font = Font(bold=cell.bold, color=cell.colour)
        for letter, width in sheet.widths.items():
            ws.column_dimensions[letter].width = width
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(os.fspath(path))


# ----- reading (the standard library)


def read(path: str | os.PathLike[str]) -> Book:
    """Every sheet's cells, fonts and widths, as sheets.py holds them."""
    try:
        with zipfile.ZipFile(path) as z:
            return _read(z)
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, ValueError, IndexError) as exc:
        raise WorkbookUnreadable(f"The workbook {Path(path).name} could not be read as an .xlsx.") from exc


def _read(z: zipfile.ZipFile) -> Book:
    names = set(z.namelist())
    strings = _strings(z) if "xl/sharedStrings.xml" in names else []
    fonts = _fonts(z) if "xl/styles.xml" in names else []
    rels = ElementTree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.iter(f"{_PKG}Relationship")}
    book = Book()
    workbook = ElementTree.fromstring(z.read("xl/workbook.xml"))
    for entry in workbook.iter(f"{_MAIN}sheet"):
        target = targets[entry.get(f"{_REL}id")]
        part = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
        book.sheets.append(_sheet(entry.get("name"), ElementTree.fromstring(z.read(part)), strings, fonts))
    return book


def _strings(z: zipfile.ZipFile) -> list[str]:
    root = ElementTree.fromstring(z.read("xl/sharedStrings.xml"))
    return ["".join(t.text or "" for t in si.iter(f"{_MAIN}t")) for si in root.iter(f"{_MAIN}si")]


def _fonts(z: zipfile.ZipFile) -> list[tuple[bool, str | None]]:
    """(bold, colour) for each cell format, by its index."""
    root = ElementTree.fromstring(z.read("xl/styles.xml"))
    fonts = []
    holder = root.find(f"{_MAIN}fonts")
    for font in (holder if holder is not None else []):
        bold = font.find(f"{_MAIN}b")
        colour = font.find(f"{_MAIN}color")
        rgb = colour.get("rgb") if colour is not None else None
        fonts.append((bold is not None and bold.get("val", "1") not in ("0", "false"),
                      rgb[-6:].upper() if rgb and rgb[-6:].upper() != "000000" else None))
    xfs = root.find(f"{_MAIN}cellXfs")
    return [fonts[int(xf.get("fontId", "0"))] if fonts else (False, None) for xf in (xfs if xfs is not None else [])]


def _sheet(title: str, root: ElementTree.Element, strings: list[str],
           fonts: list[tuple[bool, str | None]]) -> Sheet:
    sheet = Sheet(title)
    for col in root.iter(f"{_MAIN}col"):
        if col.get("width") is not None:
            for number in range(int(col.get("min")), int(col.get("max")) + 1):
                sheet.widths[column_letter(number)] = _number(col.get("width"))
    for c in root.iter(f"{_MAIN}c"):
        found = _CELL.match(c.get("r", ""))
        if found is None:
            raise ValueError("a cell with no reference")
        row, col = int(found["row"]), column_number(found["col"])
        style = int(c.get("s", "0"))
        bold, colour = fonts[style] if style < len(fonts) else (False, None)
        value = _value(c, strings)
        if value is not None or bold or colour:
            sheet.cells[(row, col)] = Cell(value, bold, colour)
    return sheet


def _value(c: ElementTree.Element, strings: list[str]) -> Any:
    formula = c.find(f"{_MAIN}f")
    if formula is not None and formula.text:
        return Formula("=" + formula.text)
    kind = c.get("t", "n")
    if kind == "inlineStr":
        return "".join(t.text or "" for t in c.iter(f"{_MAIN}t")) or None
    v = c.find(f"{_MAIN}v")
    if v is None or v.text is None:
        return None
    if kind == "s":
        return strings[int(v.text)] or None             # an empty string holds nothing, as openpyxl reads it
    if kind == "b":
        return v.text.strip() == "1"
    if kind in ("str", "e"):
        return v.text
    return _number(v.text)


def _number(text: str) -> int | float:
    value = float(text)
    return int(value) if re.fullmatch(r"-?\d+", text.strip()) else value


# ----- checking a workbook against its conversion


def check(path: str | os.PathLike[str], conv: Mapping[str, Any]) -> list[str]:
    """Every way the saved file doesn't say what conv says, in plain words; [] when it does."""
    try:
        book = read(path)
    except WorkbookUnreadable as exc:
        return [str(exc)]
    return check_book(book, conv)


def check_book(book: Book, conv: Mapping[str, Any]) -> list[str]:
    problems: list[str] = []
    want = layout(conv)
    titles = [s.title for s in book.sheets]
    if titles != [s.title for s in want.sheets]:
        return [f"The workbook's sheets are {', '.join(titles)}, not {', '.join(s.title for s in want.sheets)}."]
    try:
        evaluate(book)
    except Unevaluable as exc:
        problems.append(f"A formula could not be evaluated: {exc}.")
    problems += _totals(book, conv)
    problems += _buckets(book, conv)
    problems += _unplaced(book, conv)
    for mine, laid in zip(book.sheets, want.sheets):
        if _values(mine) != _values(laid):
            problems.append(f"{mine.title} doesn't hold the cells its conversion lays out.")
    return problems


def _values(sheet: Sheet) -> dict[tuple[int, int], Any]:
    return {k: v.value for k, v in sheet.cells.items() if v.value is not None}


def _totals(book: Book, conv: Mapping[str, Any]) -> list[str]:
    try:
        found = named(book)
    except Unevaluable:
        return []
    out = []
    parts = list(conv.get("parts") or [])
    rows = found.get("parts", [])
    if len(rows) != len(parts):
        out.append(f"{PARTS} holds {len(rows)} part rows where the conversion has {len(parts)}.")
    total = conv.get("total_board_parts")
    if found.get("board_parts") != total:
        out.append(f"{PARTS}'s total row says {_say(found.get('board_parts'))} board parts, not {total}.")
    area = sum(b.get("sqft") or 0 for b in conv.get("buckets") or [])
    if found.get("sqft") is None or abs(found["sqft"] - area) > SQFT_TOLERANCE:
        out.append(f"{PARTS}'s total row says {_say(found.get('sqft'))} SqFt, not {round(area, 2)}.")
    tape = sum(b.get("edgeband_lf") or 0 for b in conv.get("buckets") or [])
    if found.get("tape") is None or abs(found["tape"] - tape) > TAPE_TOLERANCE:
        out.append(f"{PARTS}'s total row says {_say(found.get('tape'))} ft of tape, not {round(tape, 1)}.")
    board = found.get("pulled", {}).get("board")
    if board != total:
        out.append(f"{PULLED} buckets {_say(board)} pieces as board, not {total}.")
    return out


def _buckets(book: Book, conv: Mapping[str, Any]) -> list[str]:
    found = named(book).get("buckets", {})
    out = []
    for b in conv.get("buckets") or []:
        name = f"{b.get('material')} {b.get('thk_label')}"
        mine = found.get(name)
        if mine is None:
            out.append(f"{TAKEOFF} has no {name} bucket.")
            continue
        sheets = b.get("sheets_with_waste")
        want_sheets = COULD_NOT_VERIFY if sheets is None else sheets
        if mine.get("sheets") != want_sheets:
            out.append(f"{TAKEOFF}: {name} sheets read {_say(mine.get('sheets'))}, not {want_sheets}.")
        for field, key in (("parts", "parts"), ("sqft", "sqft"), ("calc_sheets_raw", "calc_sheets_raw")):
            if mine.get(field) != b.get(key):
                out.append(f"{TAKEOFF}: {name} {field} read {_say(mine.get(field))}, not {b.get(key)}.")
    return out


def _unplaced(book: Book, conv: Mapping[str, Any]) -> list[str]:
    """Each line that could not be verified is listed as such, in order, and nowhere counted."""
    listed = list(conv.get("unplaced") or [])
    sheet = book.sheet(TAKEOFF)
    labels = [sheet.value(r, 1) for r in range(1, sheet.max_row + 1)]
    out = []
    if not listed:
        if UNPLACED_HEADER in labels:
            out.append(f"{TAKEOFF} lists lines that could not be verified, and the conversion has none.")
        return out
    if UNPLACED_HEADER not in labels:
        return [f"{TAKEOFF} doesn't list the {len(listed)} lines that could not be verified."]
    start = labels.index(UNPLACED_HEADER) + 2
    said = [sheet.value(start + i, 2) for i in range(len(listed))]
    if said != [unplaced_words(u) for u in listed]:
        out.append(f"{TAKEOFF} doesn't list the lines that could not be verified as the conversion does.")
    pulled = book.sheet(PULLED)
    marked = sum(int(pulled.value(r, 2) or 0) for r in range(2, pulled.max_row + 1)
                 if pulled.value(r, 5) == COULD_NOT_VERIFY)
    if marked != sum(int(u.get("qty") or 0) for u in listed):
        out.append(f"{PULLED} marks {marked} pieces could not verify, not"
                   f" {sum(int(u.get('qty') or 0) for u in listed)}.")
    return out


def _say(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:g}"
    return "nothing" if value is None else str(value)


